from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, ScatterChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.comments import Comment
B,O='2A78D6','EB6834'; F='Arial'
wb=Workbook()
hdr=Font(name=F,bold=True,color='FFFFFF'); hfill=PatternFill('solid',fgColor='1F3864')
inp=Font(name=F,color='0000FF'); norm=Font(name=F); bold=Font(name=F,bold=True)
thin=Side(style='thin',color='BFBFBF'); bd=Border(top=thin,bottom=thin,left=thin,right=thin)
def head(ws,r,cols):
    for j,c in enumerate(cols,1):
        x=ws.cell(r,j,c); x.font=hdr; x.fill=hfill; x.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); x.border=bd
def style(ws,r1,r2,c2):
    for r in range(r1,r2+1):
        for c in range(1,c2+1):
            x=ws.cell(r,c); x.border=bd; x.alignment=Alignment(horizontal='center')
            if x.font.color is None or x.font.color.rgb not in ('FF0000FF',): x.font=x.font if x.font.color and x.font.color.rgb=='FF0000FF' else norm
# ---------------- Data sheet (Table 2 values, original measurements)
d=wb.active; d.title='Data'
d['A1']='Table 2 – Mix proportions (per 200 g GGBS) and measured setting times (source data for Fig. 6)'; d['A1'].font=Font(name=F,bold=True,size=12)
d['A2']='Blue = measured/input values (original laboratory data, Table 2 of the manuscript; Hassan, M.Sc. thesis, GUC 2022, Table 4-1). Black = formulas.'; d['A2'].font=Font(name=F,italic=True,size=9)
cols=['Trial','Route','NaOH (g)','Na2SiO3 (g)','Extra H2O (g)','Na2O (%)','Ms','L/S','NaOH conc.','Alk/Bi','IST (min)','FST (min)','w/b = Extra H2O/200','FST − IST (min)']
head(d,4,cols)
rows=[(1,'A',110,0,0,20.5,0,0.23,'12 M',0.55,14,50),(2,'Silicate only',0,110,0,4.5,3.3,0.30,'8.15 % Na2O',0.55,None,None),
(3,'A',110,0,0,13.6,0,0.32,'8 M',0.55,33,59),(4,'A',110,0,0,17.1,0,0.27,'10 M',0.55,24,52),(5,'A',129.9,0,0,24.2,0,0.26,'12 M',0.65,44,92),
(6,'A',129.9,0,0,20.1,0,0.31,'10 M',0.65,46,77),(7,'A',129.9,0,0,16.1,0,0.37,'8 M',0.65,50,76),(8,'B',65,65,0,14.7,0.5,0.30,'12 M',0.65,10,25),
(9,'B',65,65,0,10.7,0.5,0.35,'8 M',0.65,8,20),(10,'B',93.2,46.6,0,19.2,0.3,0.30,'12 M',0.70,17,32),(11,'B',113.3,56.66,0,23.4,0.3,0.35,'12 M',0.85,24,46),
(12,'B',59.3,150.7,0,17.2,1.0,0.46,'12 M',1.05,11,23),(13,'B',44.2,165.9,0,15.0,1.3,0.47,'12 M',1.05,7,14),(14,'B',105.27,105.27,0,23.9,0.5,0.43,'12 M',1.05,18,45),
(15,'B',129.03,90.32,0,27.7,0.4,0.43,'12 M',1.10,22,62),(16,'C',17.47,44.36,64.94,6.0,1.0,0.44,'48 % Na2O',0.63,135,335),(17,'C',21.5,70.77,27.77,8.0,1.2,0.34,'48 % Na2O',0.60,None,None),
(18,'C',23.3,59.16,44.79,8.0,1.0,0.39,'48 % Na2O',0.64,15,121),(19,'C',17.51,44.38,56.6,6.0,1.0,0.42,'48 % Na2O',0.59,72,156),(20,'C',18.24,52.8,43.16,6.5,1.1,0.37,'48 % Na2O',0.57,41,146),
(21,'C',20.44,51.55,44.79,7.0,1.0,0.37,'48 % Na2O',0.58,70,150)]
for i,r in enumerate(rows):
    R=5+i
    for j,v in enumerate(r,1):
        c=d.cell(R,j,v if v is not None else '–'); c.font=inp
    d.cell(R,13,f'=E{R}/200').number_format='0.000'
    d.cell(R,14,f'=IF(ISNUMBER(L{R}),L{R}-K{R},"–")')
style(d,5,25,14)
for R in range(5,26):
    for j in range(1,13): d.cell(R,j).font=inp
d['A27']='Trials 2 and 17: no valid setting time (Trial 2 formed solid lumps on contact; Trial 17 too stiff to cast).'; d['A27'].font=Font(name=F,italic=True,size=9)
for col,w in zip('ABCDEFGHIJKLMN',[7,13,10,11,12,9,7,7,14,8,10,10,14,13]): d.column_dimensions[col].width=w
d.row_dimensions[4].height=32; d.freeze_panes='A5'
def ref(trial,col):  # cell address in Data for a given trial and column letter
    return f"=Data!{col}{4+trial}"
# ---------------- Fig 6a
a=wb.create_sheet('Fig6a_NaOH'); a['A1']='Fig. 6(a) – Route A (NaOH only): setting time vs NaOH molarity'; a['A1'].font=Font(name=F,bold=True,size=12)
head(a,3,['NaOH molarity (M)','Trial (Alk/Bi 0.55)','IST, Alk/Bi = 0.55','FST, Alk/Bi = 0.55','Trial (Alk/Bi 0.65)','IST, Alk/Bi = 0.65','FST, Alk/Bi = 0.65'])
for k,(m,t1,t2) in enumerate([(8,3,7),(10,4,6),(12,1,5)]):
    R=4+k; a.cell(R,1,m).font=inp; a.cell(R,2,t1).font=inp; a.cell(R,5,t2).font=inp
    a.cell(R,3,f'=INDEX(Data!$K$5:$K$25,MATCH(B{R},Data!$A$5:$A$25,0))'); a.cell(R,4,f'=INDEX(Data!$L$5:$L$25,MATCH(B{R},Data!$A$5:$A$25,0))')
    a.cell(R,6,f'=INDEX(Data!$K$5:$K$25,MATCH(E{R},Data!$A$5:$A$25,0))'); a.cell(R,7,f'=INDEX(Data!$L$5:$L$25,MATCH(E{R},Data!$A$5:$A$25,0))')
style(a,4,6,7)
for R in range(4,7):
    for c in (1,2,5): a.cell(R,c).font=inp
a['A8']='Change 8 → 12 M (%)'; a['A8'].font=bold
for c in 'CDFG': a[f'{c}8']=f'=({c}6-{c}4)/{c}4'; a[f'{c}8'].number_format='0%'
a['A9']='Change Alk/Bi 0.55 → 0.65 (min)'; a['A9'].font=bold
for R,lab in zip((10,11,12),(8,10,12)):
    a[f'A{R}']=f'{lab} M'; a[f'C{R}']=f'=F{R-6}-C{R-6}'; a[f'D{R}']=f'=G{R-6}-D{R-6}'
    a[f'F{R}']=f'=(F{R-6}-C{R-6})/C{R-6}'; a[f'G{R}']=f'=(G{R-6}-D{R-6})/D{R-6}'; a[f'F{R}'].number_format='0%'; a[f'G{R}'].number_format='0%'
a['B9']='ΔIST / ΔFST in C–D (min); in F–G (%)'; a['B9'].font=Font(name=F,italic=True,size=9)
for col in 'ABCDEFG': a.column_dimensions[col].width=17
a.row_dimensions[3].height=30
ch=ScatterChart(); ch.title='(a) Route A: NaOH only'; ch.style=13; ch.x_axis.title='NaOH molarity (M)'; ch.y_axis.title='Setting time (min)'
ch.y_axis.scaling.min=0; ch.y_axis.scaling.max=120; ch.x_axis.scaling.min=6; ch.x_axis.scaling.max=14; ch.x_axis.majorUnit=2
x=Reference(a,min_col=1,min_row=4,max_row=6)
for col,color,dash,sym in [(3,B,'solid','circle'),(4,O,'solid','square'),(6,B,'dash','circle'),(7,O,'dash','square')]:
    s=Series(Reference(a,min_col=col,min_row=3,max_row=6),x,title_from_data=True)
    s.marker=Marker(symbol=sym,size=7); s.marker.graphicalProperties.solidFill=color; s.marker.graphicalProperties.line.solidFill=color
    s.graphicalProperties.line.solidFill=color; s.graphicalProperties.line.width=22000; s.graphicalProperties.line.dashStyle=dash; s.smooth=False
    ch.series.append(s)
ch.x_axis.delete=False; ch.y_axis.delete=False; ch.legend.position='b'; ch.width=16; ch.height=9.5
a.add_chart(ch,'I3')
# ---------------- Fig 6b
b=wb.create_sheet('Fig6b_Ms'); b['A1']='Fig. 6(b) – Route B (Alk/Bi = 1.05): effect of silicate modulus'; b['A1'].font=Font(name=F,bold=True,size=12)
head(b,3,['Category label','Trial','Ms','IST (min)','FST (min)'])
for k,t in enumerate([14,12,13]):
    R=4+k; b.cell(R,2,t).font=inp
    b.cell(R,3,f'=INDEX(Data!$G$5:$G$25,MATCH(B{R},Data!$A$5:$A$25,0))')
    b.cell(R,1,f'=TEXT(C{R},"0.0")&" (T"&B{R}&")"')
    b.cell(R,4,f'=INDEX(Data!$K$5:$K$25,MATCH(B{R},Data!$A$5:$A$25,0))'); b.cell(R,5,f'=INDEX(Data!$L$5:$L$25,MATCH(B{R},Data!$A$5:$A$25,0))')
style(b,4,6,5)
for R in range(4,7): b.cell(R,2).font=inp
b['A8']='Change Ms 0.5 → 1.3 (%)'; b['A8'].font=bold; b['D8']='=(D6-D4)/D4'; b['E8']='=(E6-E4)/E4'; b['D8'].number_format='0%'; b['E8'].number_format='0%'
for col in 'ABCDE': b.column_dimensions[col].width=16
def bar(ws,title,xt,cats,data_min_col,max_row,ymax,anchor):
    c=BarChart(); c.type='col'; c.grouping='clustered'; c.overlap=0; c.gapWidth=80; c.title=title; c.style=10
    c.add_data(Reference(ws,min_col=data_min_col,max_col=data_min_col+1,min_row=3,max_row=max_row),titles_from_data=True)
    c.set_categories(cats)
    for s,color in zip(c.series,(B,O)):
        s.graphicalProperties.solidFill=color; s.graphicalProperties.line.solidFill=color
    c.dataLabels=DataLabelList(); c.dataLabels.showVal=True; c.dataLabels.showSerName=False; c.dataLabels.showCatName=False; c.dataLabels.showLegendKey=False; c.dataLabels.showPercent=False; c.dataLabels.position='outEnd'
    c.x_axis.title=xt; c.y_axis.title='Setting time (min)'; c.y_axis.scaling.min=0; c.y_axis.scaling.max=ymax
    c.x_axis.delete=False; c.y_axis.delete=False; c.legend.position='b'; c.width=16; c.height=9.5
    ws.add_chart(c,anchor)
bar(b,'(b) Route B: Alk/Bi = 1.05','Silicate modulus Ms',Reference(b,min_col=1,min_row=4,max_row=6),4,6,55,'G3')
# ---------------- Fig 6c
c=wb.create_sheet('Fig6c_RouteC'); c['A1']='Fig. 6(c) – Route C: effect of Na2O % and extra water (w/b)'; c['A1'].font=Font(name=F,bold=True,size=12)
head(c,3,['Category label','Trial','Na2O (%)','w/b','IST (min)','FST (min)'])
for k,t in enumerate([18,21,20,19,16]):
    R=4+k; c.cell(R,2,t).font=inp
    c.cell(R,3,f'=INDEX(Data!$F$5:$F$25,MATCH(B{R},Data!$A$5:$A$25,0))'); c.cell(R,4,f'=INDEX(Data!$M$5:$M$25,MATCH(B{R},Data!$A$5:$A$25,0))'); c.cell(R,4).number_format='0.000'
    c.cell(R,1,f'="T"&B{R}&" | "&TEXT(C{R},"0.0")&" % | "&TEXT(D{R},"0.000")')
    c.cell(R,5,f'=INDEX(Data!$K$5:$K$25,MATCH(B{R},Data!$A$5:$A$25,0))'); c.cell(R,6,f'=INDEX(Data!$L$5:$L$25,MATCH(B{R},Data!$A$5:$A$25,0))')
style(c,4,8,6)
for R in range(4,9): c.cell(R,2).font=inp; c.cell(R,4).number_format='0.000'
c['A10']='Na2O 7 % → 8 % at same w/b (T21 → T18)'; c['A10'].font=bold; c['E10']='=E4-E5'; c['F10']='=F4-F5'
c['A11']='Extra water w/b 0.283 → 0.325 (T19 → T16)'; c['A11'].font=bold; c['E11']='=E8-E7'; c['F11']='=F8-F7'
c['A12']='Values in rows 10–11 are changes in minutes.'; c['A12'].font=Font(name=F,italic=True,size=9)
c.column_dimensions['A'].width=40
for col in 'BCDEF': c.column_dimensions[col].width=12
bar(c,'(c) Route C: Na2O % and w/b','Trial | Na2O % | w/b',Reference(c,min_col=1,min_row=4,max_row=8),5,8,370,'H3')
# ---------------- Figure 6 sheet with all three charts side by side
f=wb.create_sheet('Figure 6',1); f['A1']='Figure 6 – Measured setting times by activator route (charts linked to sheets Fig6a–Fig6c)'; f['A1'].font=Font(name=F,bold=True,size=12)
f['A2']='Edit values only in the Data sheet; all tables and charts update automatically.'; f['A2'].font=Font(name=F,italic=True,size=9)
import copy
for ws,anchor in ((a,'A4'),(b,'K4'),(c,'U4')):
    ch2=copy.deepcopy(ws._charts[0]); ch2.anchor=anchor; ch2.width=15; ch2.height=9; f.add_chart(ch2,anchor)
wb.save('Figure6_Setting_Time_Trends.xlsx'); print('saved')
